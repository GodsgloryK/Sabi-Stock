"""Generate the Smart Stock project presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

GREEN = RGBColor(0x19, 0x76, 0x4F)
LIGHT_GREEN = RGBColor(0xE3, 0xEF, 0xE8)
DARK = RGBColor(0x26, 0x36, 0x2B)
GRAY = RGBColor(0x78, 0x86, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF6, 0xF8, 0xF5)
PLACEHOLDER = RGBColor(0xD8, 0xE2, 0xD8)

BRAND = "Smart Stock"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=CREAM):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, size=18, color=DARK,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=DARK, spacing=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.level = 0
    return box


def add_card(slide, left, top, width, height, title, body, icon=""):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(0xDD, 0xE5, 0xDD)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.14)
    p = tf.paragraphs[0]
    p.text = f"{icon} {title}" if icon else title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(11)
    p2.font.color.rgb = GRAY
    return shape


def add_placeholder(slide, left, top, width, height, label):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PLACEHOLDER
    shape.line.color.rgb = RGBColor(0xB7, 0xC9, 0xB8)
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p2 = tf.add_paragraph()
    p2.text = "[ screenshot ]"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(10)
    p2.font.color.rgb = GRAY
    return shape


def accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.09))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def footer(slide, page):
    add_text(slide, 0.5, 7.05, 4, 0.3, BRAND, size=10, color=GRAY)
    add_text(slide, 12.2, 7.05, 0.8, 0.3, str(page), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


# ── Slide 1: Title ──────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide, GREEN)
add_text(slide, 1, 2.2, 11.3, 1.2, BRAND, size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 1, 3.5, 11.3, 0.8, "Inventory & Sales Management Platform for Small Businesses",
         size=22, color=RGBColor(0xC9, 0xD9, 0xBD), align=PP_ALIGN.CENTER)
add_text(slide, 1, 4.6, 11.3, 0.5, "Products  ·  Inventory  ·  Sales  ·  Insights  ·  Team",
         size=14, color=RGBColor(0x9F, 0xC6, 0xA0), align=PP_ALIGN.CENTER)


# ── Slide 2: Problem ────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "The Problem", size=32, color=DARK, bold=True)
add_text(slide, 0.7, 1.05, 11, 0.5,
         "Small businesses still manage products, inventory, and sales with spreadsheets, notebooks, or disconnected systems.",
         size=16, color=GRAY)

pains = [
    ("📉", "Inaccurate stock", "Manual records drift out of sync with reality"),
    ("🔍", "No visibility", "Hard to see what sells, what's left, and what's stuck"),
    ("⚠️", "Silent stockouts", "Low-stock products go unnoticed until it's too late"),
    ("📊", "No insight", "Business performance is a guess, not a number"),
]
for i, (icon, title, body) in enumerate(pains):
    add_card(slide, 0.7 + i * 3.05, 2.0, 2.85, 1.6, title, body, icon)

add_text(slide, 0.7, 4.1, 11.9, 0.5, "The challenge: build a centralized platform that solves all four.", size=16, color=DARK, bold=True)
add_bullets(slide, 0.9, 4.7, 11.5, 2.4, [
    "Users: Business Owner  ·  Store Manager  ·  Sales Staff  ·  Administrator",
    "Core needs: product management, inventory tracking, sales recording, low-stock alerts, search & filtering, dashboard",
    "Go beyond: demand forecasting, anomaly detection, an AI assistant, reports & export, staff management",
], size=14, color=GRAY)
footer(slide, 2)


# ── Slide 3: Solution ───────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "Our Solution", size=32, color=DARK, bold=True)
add_text(slide, 0.7, 1.05, 11.9, 0.5,
         "A full-stack, multi-tenant web platform — live on Vercel (frontend) and Render (API)",
         size=16, color=GRAY)

solutions = [
    ("✅", "Product management", "Full CRUD with search, category filter, and price tracking"),
    ("✅", "Inventory tracking", "Stock auto-decrements on sale; void & restock restores it atomically"),
    ("✅", "Sales recording", "Server-computed totals & profit; live-stock product picker"),
    ("✅", "Low-stock alerts", "Configurable per-business threshold, timezone & badges"),
    ("✅", "Dashboard & analytics", "KPIs, best-seller chart, 30-day revenue/profit trend line"),
    ("✅", "Staff management", "Owner invites stock managers via single-use 24h codes; revoke anytime"),
]
for i, (icon, title, body) in enumerate(solutions):
    col = i % 3
    row = i // 3
    add_card(slide, 0.7 + col * 4.1, 1.85 + row * 1.85, 3.9, 1.6, title, body, icon)

add_text(slide, 0.7, 5.75, 11.9, 0.5, "Bonus: AI forecasting, anomaly insights & a grounded shop assistant  ·  CSV export  ·  ₦ currency  ·  timezone-correct metrics  ·  login rate-limiting",
         size=13, color=GREEN, bold=True)
footer(slide, 3)


# ── Slide 4: How it works (architecture) ────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "How It Works", size=32, color=DARK, bold=True)

layers = [
    ("Frontend", "Static HTML / CSS / vanilla JS\nDashboard · Products · Sales · Insights · Team\nChart.js for analytics & forecast charts", 0.7),
    ("API", "FastAPI (Python) on Render\nJWT auth · Argon2 hashing\nRole-based access (Owner / Stock Manager)", 4.7),
    ("Database", "PostgreSQL (Neon)\nMulti-tenant: every query scoped by business_id\nTransactional stock integrity with row locks", 8.7),
]
for title, body, left in layers:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.6), Inches(3.8), Inches(2.6)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = GREEN
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN
    for line in body.split("\n"):
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK

add_text(slide, 0.7, 4.6, 11.9, 0.5, "Security & integrity built in", size=18, color=DARK, bold=True)
add_bullets(slide, 0.9, 5.15, 11.5, 1.8, [
    "Multi-tenant isolation — membership re-verified from the database on every request; removing a member revokes access instantly",
    "Sales use FOR UPDATE row locks — no overselling under concurrent use; prices snapshotted so profit history survives price edits",
    "Products with recorded sales cannot be deleted — data integrity enforced at the database level with composite foreign keys",
], size=13, color=GRAY, spacing=6)
footer(slide, 4)


# ── Slide 5: Use of AI ──────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "Use of AI", size=32, color=DARK, bold=True)
add_text(slide, 0.7, 1.02, 11.9, 0.5,
         "AI that advises the owner — but never guesses. The model phrases facts; it never invents them.",
         size=16, color=GRAY)

ai_cards = [
    ("📈", "Demand forecasting", "Per product: sales rate, trend, days of stock, projected stockout, reorder-by date and suggested order quantity."),
    ("🔍", "Anomaly insights", "Flags sales spikes, demand drops, stale stock and below-cost selling — each with a severity."),
    ("💬", "Grounded assistant", "Chat answers built from a live snapshot: today's sales & profit, low stock, top sellers, urgent reorders."),
]
for i, (icon, title, body) in enumerate(ai_cards):
    add_card(slide, 0.7 + i * 4.1, 1.7, 3.9, 2.0, title, body, icon)

add_text(slide, 0.7, 3.95, 11.9, 0.5, "How it stays trustworthy", size=18, color=DARK, bold=True)
add_bullets(slide, 0.9, 4.5, 11.5, 2.4, [
    "The numbers are computed in pure Python from SQL history — the LLM (Gemini) only rephrases those computed facts",
    "Deterministic fallback: if the key is missing, the network fails, or the free tier rate-limits, a rules engine answers instead — the demo never fails",
    "Every response reports its source (gemini or rules), so the user always knows where the answer came from",
], size=13, color=GRAY, spacing=8)
footer(slide, 5)


# ── Slide 6: The Product (screenshots) ──────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "The Product", size=32, color=DARK, bold=True)

pages = [
    ("Dashboard", 0.7, 1.1),
    ("Insights (AI)", 7.0, 1.1),
    ("Products", 0.7, 3.15),
    ("Sales", 7.0, 3.15),
    ("Shop Assistant", 0.7, 5.2),
    ("Team & Settings", 7.0, 5.2),
]
for label, left, top in pages:
    add_placeholder(slide, left, top, 5.6, 1.85, label)
footer(slide, 6)


# ── Slide 7: Results ────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
accent_bar(slide)
add_text(slide, 0.7, 0.35, 11, 0.7, "Results", size=32, color=DARK, bold=True)

metrics = [
    ("100%", "Core requirements\ndelivered"),
    ("73", "Automated unit\ntests passing"),
    ("30", "REST API\noperations"),
    ("10", "Application\npages"),
]
for i, (num, label) in enumerate(metrics):
    left = 0.7 + i * 3.1
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.4), Inches(2.85), Inches(1.9)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(0xDD, 0xE5, 0xDD)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = num
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p2 = tf.add_paragraph()
    p2.text = label
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(12)
    p2.font.color.rgb = GRAY

add_text(slide, 0.7, 3.7, 11.9, 0.5, "Deployed & production-ready", size=18, color=DARK, bold=True)
add_bullets(slide, 0.9, 4.25, 11.5, 2.5, [
    "Live frontend on Vercel  ·  FastAPI on Render  ·  Database on Neon (PostgreSQL)",
    "AI: forecasting + anomaly insights + a grounded assistant (Gemini, with a deterministic fallback)",
    "CI runs the full 73-test suite on every push via GitHub Actions",
    "Migrations 001–006 applied; all ten pages validated end-to-end against the live API",
], size=14, color=GRAY, spacing=8)
footer(slide, 7)


prs.save("Smart-Stock-Presentation.pptx")
print("saved: Smart-Stock-Presentation.pptx")
